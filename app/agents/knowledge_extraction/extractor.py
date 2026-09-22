import fitz  # this is PyMuPDF's import name
import docx
from pptx import Presentation
import os
import pytesseract
from PIL import Image

def extract_text(file_path: str) -> str:
    text = ""
    file_path_lower = file_path.lower()
    
    if file_path_lower.endswith('.pdf'):
        doc = fitz.open(file_path)
        for page in doc:
            page_text = page.get_text()
            if not page_text.strip():
                # Fallback to OCR if no text found (e.g. scanned PDF)
                pix = page.get_pixmap()
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                page_text = pytesseract.image_to_string(img)
            text += page_text + "\n"
        doc.close()
        
    elif file_path_lower.endswith(('.png', '.jpg', '.jpeg', '.tiff', '.bmp')):
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img)
    
    elif file_path_lower.endswith('.docx'):
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
            
    elif file_path_lower.endswith('.pptx'):
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
                    
    else:
        raise ValueError(f"Unsupported file format: {file_path}")
        
    return text